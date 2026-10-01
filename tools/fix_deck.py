"""Post-conversion fix-ups for the C840 deck: strip the WSQ layer convert_deck.py left behind,
rewrite the digital-attendance slide into Course Feedback (moved after References), fix notes."""
import re
from pptx import Presentation

PATH = "courseware/AI for Product Development (C840)-v1.0.pptx"
p = Presentation(PATH)
OLD_URL = "https://www.tertiarycourses.com.sg/wsq-agentic-ai-for-product-development.html"
NEW_URL = "https://www.tertiarycourses.com.sg/ai-for-product-development.html"
OLD_REPO = "https://github.com/tertiarycourses/TGS-2024045799-Agentic-AI-for-Product-Development"
NEW_REPO = "https://github.com/tertiarycourses/C840-AI-for-Product-Development"


def ptext(para):
    return "".join(r.text for r in para.runs)


def set_para(para, line):
    runs = [r for r in para.runs if r.text] or para.runs
    runs[0].text = line
    for r in para.runs:
        if r._r is not runs[0]._r:
            r.text = ""


def find(title):
    """The slide whose shape 3 (slide title) reads exactly `title`."""
    hits = [s for s in p.slides if any(sh.shape_id == 3 and sh.has_text_frame and sh.text_frame.text == title
                                       for sh in s.shapes)]
    assert len(hits) == 1, (title, len(hits))
    return hits[0]


def put(slide, sid, text):
    sh = next(sh for sh in slide.shapes if sh.shape_id == sid)
    paras = sh.text_frame.paragraphs
    assert len(paras) == 1, (sid, len(paras))
    set_para(paras[0], text)


# ---- trainer card
put(find("Dr Alfred Ang"), 22, "Industry-oriented implementation practice")

# ---- course facts: no TSC, no assessment
facts = find("Course facts and competency context")
put(facts, 3, "Course facts")
put(facts, 10, "Non-WSQ course by Tertiary Infotech Academy")
put(facts, 22, "Daily schedule: 9:30am – 5:30pm")

# ---- day 2 card
put(find("Day 2 lesson plan"), 26, "Activities 07–12: evaluation set through executive recommendation")

# ---- LMS slide
lms = find("Access the learning platform")
put(lms, 10, "Use the LMS for learner resources and course materials.")
put(lms, 22, "Confirm the course title and code before downloading files")

# ---- Course Feedback: the digital-attendance slide, rewritten in place, moved after References
fb = find("Digital attendance is mandatory")
put(fb, 3, "Course feedback")
put(fb, 10, "Tell us how the course worked for you before you leave.")
put(fb, 14, "Complete the course feedback survey")
put(fb, 18, "Tell us what worked well")
put(fb, 22, "Suggest what we can improve")
put(fb, 26, "Your certificate of completion is emailed after the course")
refs = find("References and further study")
ids = p.slides._sldIdLst
fb_id = next(e for e in ids if p.slides.get(int(e.get("id"))) is fb)
ref_id = next(e for e in ids if p.slides.get(int(e.get("id"))) is refs)
ids.remove(fb_id)
ref_id.addnext(fb_id)

# ---- speaker notes: source URLs / code / title
for s in p.slides:
    if not s.has_notes_slide:
        continue
    for para in s.notes_slide.notes_text_frame.paragraphs:
        for r in para.runs:
            t = (r.text.replace(OLD_URL, NEW_URL).replace(OLD_REPO, NEW_REPO)
                 .replace("TGS-2024045799", "C840")
                 .replace("Agentic AI for Product Development", "AI for Product Development"))
            if t != r.text:
                r.text = t

# ---- hyperlink targets live in the slide rels, not the text
for s in p.slides:
    for rel in s.part.rels.values():
        if rel.is_external and (OLD_URL in rel.target_ref or OLD_REPO in rel.target_ref):
            new = rel.target_ref.replace(OLD_URL, NEW_URL).replace(OLD_REPO, NEW_REPO)
            rel._target = new
            rel.__dict__["target_ref"] = new

# ---- thank-you credit: convert_deck split its line break into a second, un-centred paragraph;
#      restore the source txBody and swap the code/version/date inside its single run
import copy
src = Presentation("source-wsq/courseware/PPT-Agentic AI for Product Development-v7.0.pptx")
src_tb = next(sh for sh in src.slides[-1].shapes if sh.shape_id == 20)._element.txBody
dst = next(sh for sh in p.slides[-1].shapes if sh.shape_id == 20)._element
new_tb = copy.deepcopy(src_tb)
for t in new_tb.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"):
    t.text = t.text.replace("TGS-2024045799  •  v7.0  •  11 September 2026", "C840  •  v1.0  •  1 October 2026")
dst.replace(dst.txBody, new_tb)

# ---- renumber the inline "|  Slide N" footers after the move
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if re.search(r"\|\s+Slide \d+$", r.text):
                        r.text = re.sub(r"Slide \d+$", f"Slide {i}", r.text)

p.core_properties.title = "AI for Product Development"
p.core_properties.subject = "C840"
p.save(PATH)
print("saved", len(p.slides), "slides")
